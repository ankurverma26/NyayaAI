"""
backend/ingestion/document_parser.py
─────────────────────────────────────
Document parsing for NyayaAI.
Supports extracting text and page metadata from:
- PDF (.pdf via pypdf)
- Microsoft Word (.docx via python-docx)
- Plain text / Markdown (.txt, .md)

Computes SHA-256 hash and retains page markers for contract viewer navigation.
"""
from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Union

import docx
from pypdf import PdfReader


@dataclass
class PageText:
    page_number: int  # 1-indexed
    text: str


@dataclass
class ParsedDocument:
    filename: str
    file_type: str  # "pdf" | "docx" | "txt"
    file_hash: str
    full_text: str
    pages: list[PageText] = field(default_factory=list)


def _compute_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext in (".docx", ".doc"):
        return "docx"
    return "txt"


def parse_pdf(stream: BinaryIO, filename: str = "document.pdf") -> ParsedDocument:
    """Extract text page-by-page from a PDF byte stream."""
    stream.seek(0)
    raw_bytes = stream.read()
    stream.seek(0)

    reader = PdfReader(stream)
    pages: list[PageText] = []
    full_text_parts: list[str] = []

    for idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        cleaned = page_text.strip()
        pages.append(PageText(page_number=idx, text=cleaned))
        if cleaned:
            full_text_parts.append(cleaned)

    full_text = "\n\n".join(full_text_parts)
    return ParsedDocument(
        filename=filename,
        file_type="pdf",
        file_hash=_compute_hash(raw_bytes),
        full_text=full_text,
        pages=pages,
    )


def parse_docx(stream: BinaryIO, filename: str = "document.docx") -> ParsedDocument:
    """Extract paragraphs and tables from a DOCX byte stream."""
    stream.seek(0)
    raw_bytes = stream.read()
    stream.seek(0)

    doc = docx.Document(stream)
    paragraphs: list[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs.append(t)

    # Also extract any text inside tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs.append(" | ".join(row_cells))

    full_text = "\n\n".join(paragraphs)
    # DOCX does not expose paginated layout without Word rendering engine; default to 1 page
    pages = [PageText(page_number=1, text=full_text)]

    return ParsedDocument(
        filename=filename,
        file_type="docx",
        file_hash=_compute_hash(raw_bytes),
        full_text=full_text,
        pages=pages,
    )


def parse_txt(stream: BinaryIO, filename: str = "document.txt") -> ParsedDocument:
    """Parse plain text or Markdown byte stream with UTF-8 / latin-1 fallback."""
    stream.seek(0)
    raw_bytes = stream.read()

    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("latin-1", errors="replace")

    clean_text = text.strip()
    pages = [PageText(page_number=1, text=clean_text)]

    return ParsedDocument(
        filename=filename,
        file_type="txt",
        file_hash=_compute_hash(raw_bytes),
        full_text=clean_text,
        pages=pages,
    )


def parse_document(
    source: Union[Path, str, bytes, BinaryIO], filename: str = "document.txt"
) -> ParsedDocument:
    """
    Unified entry point for document parsing.
    Accepts a filepath, bytes, or file-like binary stream.
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = path.name
        raw = path.read_bytes()
        stream = io.BytesIO(raw)
    elif isinstance(source, bytes):
        raw = source
        stream = io.BytesIO(raw)
    else:
        stream = source
        stream.seek(0)
        raw = stream.read()
        stream.seek(0)

    file_type = _detect_file_type(filename)

    if file_type == "pdf":
        return parse_pdf(stream, filename=filename)
    elif file_type == "docx":
        return parse_docx(stream, filename=filename)
    else:
        return parse_txt(stream, filename=filename)

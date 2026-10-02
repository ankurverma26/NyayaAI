"""
tests/test_ingestion.py
───────────────────────
Unit and integration tests for:
- Document parser (TXT, DOCX, PDF)
- Clause splitter (heuristics, regex, fallback)
- Ingestion service (end-to-end DB persistence)
"""
from __future__ import annotations

import io
from pathlib import Path

# pyrefly: ignore [missing-import]
import docx
# pyrefly: ignore [missing-import]
import pytest
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from backend.database.models import Base, Clause, Contract, Document
from backend.ingestion.clause_splitter import (
    classify_clause_type,
    split_into_clauses,
)
from backend.ingestion.document_parser import (
    ParsedDocument,
    parse_document,
    parse_docx,
    parse_txt,
)
from backend.ingestion.service import _infer_doc_type, ingest_contract


# ── Fixtures ───────────────────────────────────────────────────────────────────

@pytest.fixture
def sqlite_engine():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    return engine


@pytest.fixture
async def session(sqlite_engine):
    async with sqlite_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(sqlite_engine, expire_on_commit=False)
    async with factory() as s:
        yield s
    await sqlite_engine.dispose()


# ── Document Parser Tests ──────────────────────────────────────────────────────

def test_parse_txt_basic():
    content = b"This is a test contract.\nClause 1: Term."
    stream = io.BytesIO(content)
    parsed = parse_txt(stream, filename="agreement.txt")
    assert parsed.file_type == "txt"
    assert "This is a test contract." in parsed.full_text
    assert len(parsed.pages) == 1
    assert len(parsed.file_hash) == 64  # SHA-256


def test_parse_docx_in_memory():
    # Build a small in-memory docx
    doc = docx.Document()
    doc.add_paragraph("Master Services Agreement")
    doc.add_paragraph("Clause 1: Scope of Work")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Party A"
    table.rows[0].cells[1].text = "Party B"

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    parsed = parse_docx(buf, filename="msa.docx")
    assert parsed.file_type == "docx"
    assert "Master Services Agreement" in parsed.full_text
    assert "Clause 1: Scope of Work" in parsed.full_text
    assert "Party A | Party B" in parsed.full_text


def test_parse_pdf_in_memory():
    # Build a simple blank PDF with pypdf
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    buf.seek(0)

    parsed = parse_document(buf, filename="test.pdf")
    assert parsed.file_type == "pdf"
    assert len(parsed.pages) == 1


# ── Clause Splitter Tests ──────────────────────────────────────────────────────

def test_split_standard_clauses():
    text = (
        "Clause 1. Definitions and Interpretation\n"
        "In this Agreement, words have defined meanings.\n\n"
        "Clause 2. Non-Competition\n"
        "The Employee shall not compete for 2 years.\n\n"
        "Clause 3. Indemnification\n"
        "Employee will defend and indemnify Employer."
    )
    doc = ParsedDocument(
        filename="test.txt",
        file_type="txt",
        file_hash="dummy",
        full_text=text,
    )
    clauses = split_into_clauses(doc)
    assert len(clauses) == 3
    assert clauses[0].clause_number == 1
    assert "Definitions" in (clauses[0].heading or "")
    assert clauses[1].clause_type == "non_compete"
    assert clauses[2].clause_type == "indemnity"


def test_classify_clause_types():
    assert classify_clause_type("Non-Compete", "restraint of trade") == "non_compete"
    assert classify_clause_type("Indemnity", "defend and hold harmless") == "indemnity"
    assert classify_clause_type("Liability", "limitation of liability cap") == "liability"
    assert classify_clause_type("Arbitration", "Arbitration and Conciliation Act") == "arbitration"
    assert classify_clause_type("Governing Law", "courts of Mumbai") == "jurisdiction"
    assert classify_clause_type("Random Clause", "general business terms") == "general"


def test_split_fallback_paragraphs():
    # Plain text without numbered clause headings
    text = (
        "This is an informal memo between parties.\n"
        "The first term is that work starts tomorrow morning.\n\n"
        "The second term is that confidentiality will be preserved.\n"
        "Neither party shall share proprietary secrets with anyone."
    )
    doc = ParsedDocument(
        filename="memo.txt",
        file_type="txt",
        file_hash="dummy",
        full_text=text,
    )
    clauses = split_into_clauses(doc)
    assert len(clauses) >= 2
    assert any(c.clause_type == "confidentiality" for c in clauses)


# ── Ingestion Service Integration Tests ────────────────────────────────────────

@pytest.mark.asyncio
async def test_ingest_sample_contract(session: AsyncSession):
    sample_path = Path("data/sample_contracts/standard_employment_agreement.txt")
    assert sample_path.exists()

    doc, contract, clauses = await ingest_contract(
        session=session,
        content=sample_path,
        filename=sample_path.name,
    )

    assert doc.id is not None
    assert doc.file_type == "txt"
    assert len(doc.file_hash) == 64

    assert contract.id is not None
    assert contract.doc_type == "employment"
    assert contract.document_id == doc.id

    assert len(clauses) >= 9

    # Verify queryable from DB
    loaded_clauses = (
        (await session.execute(select(Clause).where(Clause.contract_id == contract.id)))
        .scalars()
        .all()
    )
    assert len(loaded_clauses) == len(clauses)

    types = [c.clause_type for c in loaded_clauses]
    assert "non_compete" in types
    assert "indemnity" in types
    assert "liability" in types
    assert "arbitration" in types
    assert "jurisdiction" in types


def test_infer_doc_type():
    assert _infer_doc_type("Employment Agreement", "salary and bonus") == "employment"
    assert _infer_doc_type("Non-Disclosure Agreement", "confidential secrets") == "nda"
    assert _infer_doc_type("Master Services Agreement", "consultancy services") == "service"
    assert _infer_doc_type("Lease Deed", "tenancy monthly rent") == "lease"
